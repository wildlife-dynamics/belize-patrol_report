"""
Builds resources/templates/belize_patrol_report_template.docx, a docxtpl
Jinja2 .docx template for the Belize Fisheries Department patrol & vessel
monitoring report - the production counterpart to .scripts/build_docx_report.py
(which directly populates a docx for prototyping; this one is the real
Jinja2 template ecoscope_workflows_ext_belize.tasks.generate_belize_patrol_report
renders via docxtpl on every workflow run).

Section layout matches the source PDF (SWCMR_Report_000054.pdf), per the PRD:
title/metadata block, Patrol Effort Summary, Staff Effort, Patrols by Mandate
chart, Type of Vessels Inspected chart, Vessels Documented, Recreational
Fishing Vessels, Recreational Tourism Vessels, Commercial Fishing Vessels
Encountered, Fishers Documented, Map of Activities Documented, plus one
addition beyond the source PDF: Photos (event attachments - inspection/catch
photos - downloaded via download_event_attachments). No Joint Patrol Team
section (no matching event type - PRD open question #4). No Site/Station
Name (dropped - PRD open question #6, no site/station field exists anywhere
in the schema, and this account covers more than one protected area).

Context keys match generate_belize_patrol_report's context dict exactly
(ecoscope_workflows_ext_belize/tasks/_report.py) - a table's row-loop reads
{{ row.<safe_key(column)> }} for each of its own fixed columns, matching
whatever assemble_patrol_summary / flatten_commercial_fishing_vessels /
explode_commercial_fishing_fishers / filter_recreational_vessels_by_category
actually output.

Row-loop tables MUST have their "{%tr for %}"/"{%tr endfor %}" markers each
alone in their own table row - docxtpl's regex-based preprocessor binds a
"{% %}"-shaped tag to the LAST such tag found within the enclosing <w:tr>, so
putting a for-tag and endfor-tag in the SAME row as the data collapses the
whole row down to just the endfor, silently discarding everything between
(confirmed in curacao-nest_hatching/.scripts/build_report_template.py).

Run with:
  cd /Users/zak/Documents/w-dynamics/belize-patrol_report/ecoscope-workflows-belize-patrol-report-workflow
  pixi run python ../.scripts/build_report_template.py
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

OUT_PATH = Path(__file__).parent.parent / "resources" / "templates" / "belize_patrol_report_template.docx"

HEADING_FONT = "Georgia"
BODY_FONT = "Calibri"
BLACK = RGBColor(0, 0, 0)
HEADER_FILL_HEX = "B7DFCB"  # light teal/mint, matching the source PDF's table header shading


def _style_run(run, font=None, size=None, bold=None, color=BLACK):
    if font:
        run.font.name = font
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = rPr.makeelement(qn("w:rFonts"), {})
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), font)
        rFonts.set(qn("w:hAnsi"), font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
    tcPr.append(shd)


def _safe_key(h: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]", "_", h)
    return re.sub(r"_+", "_", s).strip("_")


def _heading(doc, text, size=16):
    p = doc.add_paragraph()
    _style_run(p.add_run(text), font=HEADING_FONT, size=size, bold=True)
    return p


def _row_loop_table(doc, columns, loop_var):
    """A docxtpl row-loop table: header row, then a bare for-marker row, one
    data row (the template row - this is what repeats), then a bare endfor
    marker row. See this file's docstring for why the for/endfor markers must
    each be alone in their own row."""
    table = doc.add_table(rows=4, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate(columns):
        header_cell = table.rows[0].cells[i]
        _shade_cell(header_cell, HEADER_FILL_HEX)
        p = header_cell.paragraphs[0]
        _style_run(p.add_run(col), font=BODY_FONT, size=9, bold=True)

        data_cell = table.rows[2].cells[i]
        p = data_cell.paragraphs[0]
        _style_run(p.add_run(f"{{{{ row.{_safe_key(col)} }}}}"), font=BODY_FONT, size=9)
        # `col` is now already the exact display_name/column string the report's
        # spec.yaml (or ecoscope_workflows_ext_belize task) uses - matching that
        # verbatim keeps this header identical to the same table's HTML dashboard
        # widget, and _safe_key(col) must derive the SAME jinja variable name that
        # _df_to_rows (ecoscope_workflows_ext_belize/tasks/_report.py) computes
        # from that same string - renaming a column anywhere upstream (spec.yaml
        # display_name, or a custom task's output columns) means this script's
        # `columns` lists below must be updated too, and the template rebuilt,
        # or the row loop still runs but every cell renders blank (confirmed the
        # hard way - see session history).

    table.rows[1].cells[0].text = f"{{%tr for row in {loop_var}.rows %}}"
    table.rows[3].cells[0].text = "{%tr endfor %}"
    return table


def _image_placeholder(doc, var_name):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{{{{ {var_name} }}}}")


def _photo_gallery_table(doc, loop_var):
    """Same row-loop shape as _row_loop_table, but for a list of
    {"image": InlineImage, "info": str} dicts (generate_belize_patrol_report's
    own photos context, from download_event_attachments) rather than a
    dataframe's rows - so the Jinja keys are the dict's own "image"/"info",
    not _safe_key(<column display name>)."""
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (header, key) in enumerate([("Photo", "image"), ("Details", "info")]):
        header_cell = table.rows[0].cells[i]
        _shade_cell(header_cell, HEADER_FILL_HEX)
        p = header_cell.paragraphs[0]
        _style_run(p.add_run(header), font=BODY_FONT, size=9, bold=True)

        data_cell = table.rows[2].cells[i]
        p = data_cell.paragraphs[0]
        _style_run(p.add_run(f"{{{{ photo.{key} }}}}"), font=BODY_FONT, size=9)

    table.rows[1].cells[0].text = f"{{%tr for photo in {loop_var} %}}"
    table.rows[3].cells[0].text = "{%tr endfor %}"
    return table


doc = Document()

# --- Title / metadata block ---
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
_style_run(title.add_run("Belize Fisheries Department"), font=HEADING_FONT, size=24, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
_style_run(subtitle.add_run("Patrol & Vessel Monitoring Report"), font=HEADING_FONT, size=16)

meta1 = doc.add_paragraph()
meta1.alignment = WD_ALIGN_PARAGRAPH.CENTER
_style_run(meta1.add_run("Report Generated by: "), font=BODY_FONT, size=11)
_style_run(meta1.add_run("{{ generated_by }}"), font=BODY_FONT, size=11, bold=True)
_style_run(meta1.add_run("    Report Generated on: "), font=BODY_FONT, size=11)
_style_run(meta1.add_run("{{ generated_on }}"), font=BODY_FONT, size=11, bold=True)

meta2 = doc.add_paragraph()
meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
_style_run(meta2.add_run("Report Start Date: "), font=BODY_FONT, size=11)
_style_run(meta2.add_run("{{ report_start_date }}"), font=BODY_FONT, size=11, bold=True)
_style_run(meta2.add_run("    Report End Date: "), font=BODY_FONT, size=11)
_style_run(meta2.add_run("{{ report_end_date }}"), font=BODY_FONT, size=11, bold=True)

doc.add_paragraph()

# --- Patrol Effort Summary ---
_heading(doc, "Patrol Effort Summary")
_row_loop_table(doc, ["Number of Patrols", "Distance (km)", "Number of Nights", "Number of Patrol Hours", "Fuel Used (gal)"], "patrol_effort_summary")

# --- Staff Effort ---
_heading(doc, "Staff Effort")
_row_loop_table(doc, ["Officer Name", "Number of Patrols", "Number of Nights", "Distance (km)", "Number of Hours"], "staff_effort")

# --- Patrols by Mandate (chart only, no table - matches source PDF) ---
_heading(doc, "Patrols by Mandate")
_image_placeholder(doc, "patrols_by_mandate_chart")

# --- Type of Vessels Inspected (chart only) ---
_heading(doc, "Type of Vessels Inspected")
_image_placeholder(doc, "vessels_inspected_chart")

# --- Vessels Documented ---
_heading(doc, "Vessels Documented")
_row_loop_table(doc, ["Recreational Fishing", "Commercial Fishing", "Subsistence Fishing", "Tourism Vessel"], "vessels_documented")

# --- Recreational Fishing Vessels ---
_heading(doc, "Recreational Fishing Vessels")
_row_loop_table(doc, ["Vessel Name", "Vessel Registration Number", "Number of Passengers", "Type of Activities"], "recreational_fishing_vessels")

# --- Recreational Tourism Vessels ---
_heading(doc, "Recreational Tourism Vessels")
_row_loop_table(doc, ["Vessel Name", "Vessel Registration Number", "Number of Passengers", "Type of Activities"], "recreational_tourism_vessels")

# --- Commercial Fishing Vessels Encountered ---
_heading(doc, "Commercial Fishing Vessels Encountered")
_row_loop_table(
    doc,
    ["Vessel Name", "Vessel Type", "Vessel home port", "Vessel License/ Registration Number", "Name of Captain", "Number of Crew", "Type of Inspection", "Type of Gear", "Number of Infractions"],
    "commercial_fishing_vessels",
)

# --- Fishers Documented ---
_heading(doc, "Fishers Documented")
_row_loop_table(doc, ["Vessel Name", "Name of Fisher", "License Number"], "fishers_documented")

# --- Map of Activities Documented ---
_heading(doc, "Map of Activities Documented")
_image_placeholder(doc, "activities_map")

# --- Photos (event attachments - inspection/catch photos, etc.) ---
_heading(doc, "Photos")
_photo_gallery_table(doc, "photos")

OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT_PATH)
print(f"Wrote: {OUT_PATH}")
